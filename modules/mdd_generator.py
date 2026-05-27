import streamlit as st
import pandas as pd
import re
import os
from io import BytesIO
from openpyxl import load_workbook
from openpyxl.cell.text import Text
import streamlit.components.v1 as components
from modules.output_utils import escape_dimensions_string, javascript_string_literal

def copy_to_clipboard(text):
    serialized_text = javascript_string_literal(text)
    copy_js = f"""
        <script>
        function copyText() {{
            const text = {serialized_text};
            navigator.clipboard.writeText(text).then(() => {{
                alert('코드가 클립보드에 복사되었습니다!');
            }});
        }}
        </script>
        <button class="copy-btn" onclick="copyText()" style="
            background-color: #f0f2f6; /* 배경색 */
            color: #31333f;            /* 글자색 */
            border: 1px solid rgba(49, 51, 63, 0.2); /* 회색 테두리 */
            padding: 0px 20px;
            border-radius: 8px;
            cursor: pointer;
            font-weight: 400;
            font-size: 1rem;
            width: 100%;
            height: 45px;
            transition: background-color 0.2s;
        ">📋 코드 전체 복사</button>
    """
    components.html(copy_js, height=60)

def get_rich_styled_text(cell):
    """엑셀 셀의 부분 볼드(b)와 밑줄(u) 설정을 HTML 태그로 변환"""
    if cell.value is None: return ""
    if isinstance(cell.value, Text) and cell.value.richText:
        parts = []
        for segment in cell.value.richText:
            if isinstance(segment, str): parts.append(segment)
            else:
                t = segment.t
                if segment.font:
                    if segment.font.u: t = f"<u>{t}</u>"
                    if segment.font.b: t = f"<b>{t}</b>"
                parts.append(t)
        return "".join(parts).replace("\n", " ")
    text = str(cell.value).replace("\n", " ")
    if hasattr(cell, "font") and cell.font:
        if cell.font.u: text = f"<u>{text}</u>"
        if cell.font.b: text = f"<b>{text}</b>"
    return text


def generate_mdd_script(df, template_dict):
    final_script = ""
    for q_num, group in df.groupby("Question Number", sort=False):
        q_type = str(group["Type"].iloc[0]).lower().strip().replace("<b>", "").replace("</b>", "")
        q_text = escape_dimensions_string(group["Question Text"].iloc[0])
        metadata_script = template_dict.get(q_type, "")
        q_id = str(q_num).replace("-", "x").replace("<b>", "").replace("</b>", "")

        unique_group = group.drop_duplicates(subset=["Attribute Number"])
        attr_count = len(unique_group)
        calculated_width = round(100 / attr_count, 1) if attr_count > 0 else 20

        if q_type in ["single2", "dynamicgrid2"]:
            if not q_text.startswith("{#rp_css}"):
                q_text = "{#rp_css}" + q_text

        if "dynamicgrid" in q_type:
            metadata_script = re.sub(
                r"columnGrid\$all\$width\s*=\s*[\d\.]+",
                f"columnGrid$all$width = {calculated_width}",
                metadata_script,
            )
        elif q_type == "single2":
            for prop in ["grid$medium$width", "grid$large$width", "grid$extralarge$width"]:
                metadata_script = re.sub(
                    prop.replace("$", r"\$") + r"\s*=\s*[\d\.]+",
                    f"{prop} = {calculated_width}",
                    metadata_script,
                )

        items = []
        for _, row in unique_group.iterrows():
            v = str(row["Attribute Number"]).split(".")[0].replace("<b>", "").replace("</b>", "")
            lbl = escape_dimensions_string(row["Attribute Text"])

            if not v or v == "nan" or v == "":
                continue

            if q_type in ["single2", "dynamicgrid2"]:
                lbl = f"<p>{v}</p>{lbl}"

            suffix = " fix exclusive" if "exe" in str(row["Attribute exe"]).lower() else ""
            items.append(f'        _{v}{"@" if suffix else ""} "{lbl}"{suffix}')

        items_str = ",\n".join(items)

        if q_type == "rank1":
            block = f"""    {q_id} "{q_text}"
{metadata_script}
    loop
    {{
{items_str}
    }} ran fields -
    (
        slice ""
        long
        defaultanswer(0 );
    ) expand grid;

    {q_id}xFin ""
    loop
    {{
        _1 "1",
        _2 "2",
        _3 "3"
    }} fields -
    (
        slice ""
        categorical [1..1]
        {{
{items_str}
        }};
    ) column expand;\n\n"""
        elif q_type == "rank2":
            if not q_text.startswith("{#Fnstr}"):
                q_text = "{#Fnstr}" + q_text

            block = f"""    {q_id}block - block fields
    (
        {q_id} "{q_text}"
{metadata_script}
        loop
        {{
{items_str}
        }} ran fields -
        (
            slice ""
            long [0 .. 3]
            defaultanswer(0 );
        ) expand grid;

        {q_id}xother "Hidden 으로 가려질 other Text 자리"
            style(
                Control(
                    Type = "SingleLineEdit"
                )
            )
        text;
    );

    {q_id}xFin ""
    loop
    {{
        _1 "1",
        _2 "2",
        _3 "3"
    }} fields -
    (
        slice ""
        categorical [0..1]
        {{
{items_str}
        }};
    ) column expand;\n\n"""
        elif "dynamicgrid" in q_type:
            block = f"""    {q_id} "{q_text}"
{metadata_script}
    loop
    {{
        use List -
    }} ran fields -
    (
        slice ""
        categorical [1..1]
        {{
{items_str}
        }};
    ) expand grid;\n\n"""
        else:
            range_val = "[1..]" if "multi" in q_type else "[1..1]"
            block = f'    {q_id} "{q_text}"\n{metadata_script}\n    categorical {range_val}\n    {{\n{items_str}\n    }};\n\n'

        final_script += block

    return final_script

def mdd_generator_page():
    st.markdown("<p style='color: black; font-size: 0.9rem; margin-bottom: -10px;'>ScriptCoded.xlsx 문서 형식에 맞춰서 업로드하세요.</p>", unsafe_allow_html=True)

    # 예시 파일 다운로드 버튼 추가
    # 현재 파일(modules/mdd_generator.py) 위치를 기준으로 template 폴더 경로 계산
    current_dir = os.path.dirname(os.path.abspath(__file__))
    template_path = os.path.join(os.path.dirname(current_dir), "template", "ScriptCoded.xlsx")

    try:
        with open(template_path, "rb") as f:
            st.download_button(
                label="📁 ScriptCoded.xlsx 예시 파일 다운로드",
                data=f,
                file_name="ScriptCoded.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
    except FileNotFoundError:
        st.warning(f"예시 파일({template_path})을 찾을 수 없습니다.")

    uploaded_survey = st.file_uploader(
        label="설문지 업로드 창",
        label_visibility="collapsed",
        type=["xlsx"]
    )

    if uploaded_survey:
        info_df = pd.read_excel(uploaded_survey, sheet_name='info', header=None)
        template_dict = {str(row[0]).lower().strip(): str(row[1]) for _, row in info_df.iterrows()}

        wb = load_workbook(uploaded_survey, data_only=False)
        ws = wb['Question']
        headers = [str(cell.value).strip() if cell.value else "" for cell in ws[1]]

        full_data = []
        l_type, l_num, l_text = "", "", ""

        for row in ws.iter_rows(min_row=2):
            curr = {}
            for i, cell in enumerate(row):
                if i < len(headers):
                    col = headers[i]
                    val = get_rich_styled_text(cell)
                    if col == "Type":
                        if val and not val.isspace(): l_type = val
                        curr[col] = l_type
                    elif col == "Question Number":
                        if val and not val.isspace(): l_num = val
                        curr[col] = l_num
                    elif col == "Question Text":
                        if val and not val.isspace(): l_text = val
                        curr[col] = l_text
                    else:
                        curr[col] = val
            full_data.append(curr)

        df = pd.DataFrame(full_data)

        col1, col2, _ = st.columns([1.5, 1.5, 7])
        generate_clicked = col1.button("🚀 MDD 스크립트 생성", use_container_width=True)

        if generate_clicked:
            final_script = generate_mdd_script(df, template_dict)
            
            with col2:
                copy_to_clipboard(final_script)
            
            st.code(final_script)
