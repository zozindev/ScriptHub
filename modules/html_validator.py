import re
from io import BytesIO

import streamlit as st
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill


HTML_TAG_RE = re.compile(r"<[^>]+>")
MALFORMED_CLOSING_TAG_RE = re.compile(r"</\s*[a-zA-Z0-9]+\s+>")
MALFORMED_OPENING_TAG_RE = re.compile(r"<[a-zA-Z0-9]+\s+>")
TAG_RE = re.compile(r"<(/?)([a-zA-Z0-9]+)\b[^>]*>")
VOID_TAGS = frozenset({"br", "img", "hr", "input", "link", "meta"})


def is_html_error(html_content):
    if not html_content or not isinstance(html_content, str):
        return False
    if not HTML_TAG_RE.search(html_content):
        return False
    if MALFORMED_CLOSING_TAG_RE.search(html_content):
        return True
    if MALFORMED_OPENING_TAG_RE.search(html_content):
        return True

    stack = []
    for is_closing, tag_name in TAG_RE.findall(html_content):
        tag_name = tag_name.lower()
        if tag_name in VOID_TAGS:
            continue
        if not is_closing:
            stack.append(tag_name)
        else:
            if not stack or stack[-1] != tag_name:
                return True
            stack.pop()
    return bool(stack)


def html_validator_page():
    uploaded_file = st.file_uploader("엑셀 파일 업로드", type=["xlsx"])
    if uploaded_file and st.button("검사 및 리스트 생성"):
        output = BytesIO()
        wb = load_workbook(uploaded_file, data_only=False)
        try:
            error_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            link_font = Font(color="0563C1", underline="single")
            error_log = []
            for sheet in wb.worksheets:
                if sheet.title == "TagErrList":
                    continue
                for row in sheet.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, str) and is_html_error(cell.value):
                            cell.fill = error_fill
                            error_log.append((sheet.title, cell.coordinate))

            if error_log:
                if "TagErrList" in wb.sheetnames:
                    del wb["TagErrList"]
                err_sheet = wb.create_sheet("TagErrList", 0)
                err_sheet.append(["시트명", "에러 셀 위치 (클릭 시 이동)"])
                err_sheet.column_dimensions["A"].width = 25
                err_sheet.column_dimensions["B"].width = 40
                for sheet_name, address in error_log:
                    new_row_idx = err_sheet.max_row + 1
                    err_sheet.cell(row=new_row_idx, column=1, value=sheet_name)
                    link_cell = err_sheet.cell(
                        row=new_row_idx,
                        column=2,
                        value=f"{sheet_name} 시트 - {address}",
                    )
                    link_cell.hyperlink = f"#{sheet_name}!{address}"
                    link_cell.font = link_font
                wb.save(output)
                output.seek(0)
                st.error(f"✅ 총 {len(error_log)}개의 에러 발견!")
                st.download_button("결과 파일 다운로드", output, f"checked_{uploaded_file.name}")
            else:
                st.success("🎉 모든 태그가 정상입니다!")
        finally:
            wb.close()
